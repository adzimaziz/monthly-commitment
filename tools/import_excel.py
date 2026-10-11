"""Tukar fail Excel lama kepada fail backup JSON untuk app Komitmen Bulanan.

Guna:
    python tools/import_excel.py "Monthly Commitment Adzim_2026_08.xlsx" "Monthly Commitment Wahidah.xlsx" private/backup-awal.json [mula] [akhir]

[mula]/[akhir] dalam format YYYY-MM (lalai: 2026-01 hingga 2026-08). Item yang tiada rekod dalam julat
tu tetap dimasukkan tetapi ditanda "dah tamat" (boleh guna semula dari Tambah item).

Lepas tu buka app > menu > Pulihkan backup > pilih fail JSON tu.
Fail output mengandungi data peribadi - JANGAN commit ke GitHub (folder private/ dah ada dalam .gitignore).
"""
import json
import re
import sys
from datetime import datetime

import openpyxl


def month_key(start_year, start_month, offset):
    m = start_month - 1 + offset
    return f"{start_year + m // 12:04d}-{m % 12 + 1:02d}"


def clean_note(text):
    if not text:
        return ""
    if "[Threaded comment]" in text and "Comment:" in text:
        text = text.split("Comment:", 1)[1]
    lines = [l.rstrip() for l in text.replace("\r", "").split("\n")]
    out = []
    for l in lines:
        s = l.strip()
        if s.startswith("======") or s.startswith("ID#"):
            continue
        if re.match(r"^.+\(\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}\)$", s):
            continue
        if re.match(r"^[^:]{3,40}:$", s) and not out:
            continue  # "Mohammad Adzim bin Ab Aziz:" author prefix
        out.append(s)
    return "\n".join(out).strip()


def cell_value(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2) if isinstance(v, float) else v
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    return s or None


def import_sheet(ws, person, name_col, code_col, acc_col, first_col, start, item_rows, ref_rows, first_key, last_key):
    items, months, refs = [], {}, []
    for r in item_rows:
        name = ws.cell(r, name_col).value
        acc = ws.cell(r, acc_col).value
        if name is None:
            continue
        iid = f"{person}-{r}"
        item = {"id": iid, "name": str(name).strip(), "jompay": "", "account": "", "info": ""}
        if code_col:
            code = ws.cell(r, code_col).value
            item["jompay"] = "" if code is None else str(code)
        if acc is not None:
            item["account"] = str(acc).strip()
        acc_cell = ws.cell(r, acc_col)
        if acc_cell.comment:
            item["info"] = clean_note(acc_cell.comment.text)
        items.append(item)
        for c in range(first_col, ws.max_column + 1):
            cell = ws.cell(r, c)
            key = month_key(start[0], start[1], c - first_col)
            if key > last_key:
                break
            if key < first_key:
                continue
            v = cell_value(cell.value)
            note = clean_note(cell.comment.text) if cell.comment else ""
            if v is None and not note:
                continue
            entry = {"v": "" if v is None else v, "paid": v is not None}
            if note:
                entry["note"] = note
            months.setdefault(key, {}).setdefault(person, {})[iid] = entry
    for r in ref_rows:
        label = ws.cell(r, name_col).value
        val = ws.cell(r, acc_col).value
        notes = [clean_note(ws.cell(r, c).comment.text) for c in range(1, ws.max_column + 1) if ws.cell(r, c).comment]
        if label is None and val is None and not notes:
            continue
        parts = ([str(val).strip()] if val is not None else []) + notes
        refs.append({"label": "" if label is None else str(label).strip(), "value": "\n".join(parts)})
    # Hide items with no value in the last 12 months
    recent = sorted(months)[-12:]
    for it in items:
        it["hidden"] = not any(it["id"] in months[k].get(person, {}) for k in recent)
    return items, months, refs


def merge(dst, src):
    for k, per in src.items():
        dst.setdefault(k, {}).update(per)


def main():
    adzim_path, wahidah_path, out = sys.argv[1:4]
    first_key = sys.argv[4] if len(sys.argv) > 4 else "2026-01"
    last_key = sys.argv[5] if len(sys.argv) > 5 else "2026-08"

    ws = openpyxl.load_workbook(adzim_path)["Monthly Commitment"]
    a_items, a_months, a_refs = import_sheet(
        ws, "adzim", name_col=2, code_col=3, acc_col=4, first_col=5, start=(2014, 8),
        item_rows=range(3, 45), ref_rows=list(range(47, 60)) + [63], first_key=first_key, last_key=last_key)
    # Row 63 has only a note in Mac 2025 column
    for r in a_refs:
        if not r["label"] and r["value"].startswith("No Tali"):
            r["label"] = "Nota sekolah (Mac 2025)"

    ws = openpyxl.load_workbook(wahidah_path)["Sheet1"]
    w_items, w_months, w_refs = import_sheet(
        ws, "wahidah", name_col=3, code_col=None, acc_col=4, first_col=5, start=(2021, 4),
        item_rows=range(5, 18), ref_rows=[19], first_key=first_key, last_key=last_key)
    for r in w_refs:
        if not r["label"]:
            r["label"] = "Akaun bank"

    months = {}
    merge(months, a_months)
    merge(months, w_months)
    data = {
        "app": "monthly-commitment",
        "version": 1,
        "people": [{"id": "adzim", "name": "Adzim"}, {"id": "wahidah", "name": "Wahidah"}],
        "items": {"adzim": a_items, "wahidah": w_items},
        "months": dict(sorted(months.items())),
        "refs": {"adzim": a_refs, "wahidah": w_refs},
        "updatedAt": datetime.now().isoformat(timespec="seconds"),
    }
    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    n_entries = sum(len(p) for m in months.values() for p in m.values())
    n_notes = sum(1 for m in months.values() for p in m.values() for e in p.values() if e.get("note"))
    print(f"{len(months)} bulan, {len(a_items)}+{len(w_items)} item, {n_entries} sel, {n_notes} nota -> {out}")


if __name__ == "__main__":
    main()
